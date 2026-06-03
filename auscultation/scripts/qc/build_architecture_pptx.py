#!/usr/bin/env python3
"""把模型结构图生成为【原生可编辑】的 .pptx(每个模块=圆角矩形+文字, 箭头=连接线)。
传到 Google Drive 用 Google Slides 打开即可逐元素编辑; 也可导入 PowerPoint / Figma。
与 docs/model_architecture.md / draw_architecture.py 一致。
输出 docs/figures/model_architecture.pptx。用法: python scripts/qc/build_architecture_pptx.py
"""
import os

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
OUT = os.path.join(ROOT, "docs", "figures", "model_architecture.pptx")

# 16:9 画布
SW, SH = 13.333, 7.5
MX, MY = 0.25, 0.25
UW, UH = SW - 2 * MX, SH - 2 * MY

# 配色(RGB)
def C(h): return RGBColor.from_string(h)
COL = {
    "in": ("ECEFF3", "8A93A0"), "per": ("DBE7FF", "2F5FD0"), "if": ("FFE7A3", "C2901A"),
    "vla": ("D5F0DC", "2E8B57"), "aud": ("FBD9C4", "D9722B"), "rob": ("DDE3EA", "5B6675"),
    "white": ("FFFFFF", "333333"),
}
TCOL = {"per": "1A3A8A", "if": "7A5500", "vla": "1D5E38", "aud": "8A4416", "rob": "3A4250", "in": "333333"}


# 坐标变换: 自定义 0..100(x右, y上) → pptx Inches(y下)
def PX(x): return Inches(MX + x / 100.0 * UW)
def PY(y): return Inches(MY + (100 - y) / 100.0 * UH)
def PW(w): return Inches(w / 100.0 * UW)
def PH(h): return Inches(h / 100.0 * UH)


def rrect(slide, x, y, w, h, fill, line, title, body=None, tsize=16, bsize=11,
          tcol="111111", title_top=False, bold=True):
    """圆角矩形 + 文字。(x,y)=左下角(自定义系)。"""
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PX(x), PY(y + h), PW(w), PH(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = C(fill)
    sp.line.color.rgb = C(line); sp.line.width = Pt(1.6)
    sp.shadow.inherit = False
    tf = sp.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP if (title_top or body) else MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = title
    r.font.size = Pt(tsize); r.font.bold = bold; r.font.color.rgb = C(tcol)
    if body:
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER; p2.space_before = Pt(4)
        rb = p2.add_run(); rb.text = body
        rb.font.size = Pt(bsize); rb.font.color.rgb = C(tcol)
    return sp


def arrow(slide, x1, y1, x2, y2, line, width=2.0, dashed=False):
    cxn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, PX(x1), PY(y1), PX(x2), PY(y2))
    cxn.line.color.rgb = C(line); cxn.line.width = Pt(width)
    ln = cxn.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    if dashed:
        d = ln.makeelement(qn("a:prstDash"), {"val": "dash"}); ln.insert(0, d)
    return cxn


def textlabel(slide, x, y, t, size=10, col="333333", w=20, h=4):
    tb = slide.shapes.add_textbox(PX(x - w / 2), PY(y + h / 2), PW(w), PH(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = t; r.font.size = Pt(size); r.font.color.rgb = C(col)
    return tb


def main():
    prs = Presentation()
    prs.slide_width = Inches(SW); prs.slide_height = Inches(SH)
    slide = prs.slides.add_slide(prs.slide_layouts[6])   # 空白

    # 标题
    textlabel(slide, 50, 97.5, "Robotic Auscultation — Perception + Action Architecture",
              size=22, col="111111", w=90, h=5)

    # 容器(先画, 在底层)
    rrect(slide, 2, 20, 16, 68, *COL["in"], "Inputs", tsize=16, tcol=TCOL["in"], title_top=True)
    rrect(slide, 23, 57, 42, 31, *COL["per"], "①  Perception   (RGB → dense IUV)",
          tsize=16, tcol=TCOL["per"], title_top=True)
    rrect(slide, 70, 40, 28, 48, *COL["vla"], "②  Arm-Control Policy  (VLA, π0)",
          tsize=16, tcol=TCOL["vla"], title_top=True)

    # 输入子项
    for name, yy in [("Third-person RGB-D", 75), ("Wrist RGB", 65), ("Proprioception", 55),
                     ("Force / Torque", 45), ("Audio", 35), ("Language", 25)]:
        rrect(slide, 4, yy, 12, 8, *COL["white"], name, tsize=11, bold=False)

    # 感知子模块
    rrect(slide, 25, 67, 12, 12, "FFFFFF", COL["per"][1], "Vision\nEncoder", tsize=12)
    rrect(slide, 39, 67, 11, 12, "FFFFFF", COL["per"][1], "Dense IUV\nHead", tsize=12)
    rrect(slide, 52, 67, 11, 12, "FFFFFF", COL["per"][1], "Dense IUV\n(I, U, V)", tsize=12)
    rrect(slide, 25, 58.6, 38, 6.6, "EEF4FF", COL["per"][1],
          "Localize task landmarks  (Heart → cardiac · Lung → pulmonary)", tsize=11, bold=False)
    arrow(slide, 37, 73, 39, 73, COL["per"][1]); arrow(slide, 50, 73, 52, 73, COL["per"][1])

    # 目标交接 接口
    rrect(slide, 23, 44.5, 42, 10, *COL["if"], "Target Hand-off   (Perception → Action)",
          body='"where to listen":  target heatmap  +  3D target point (base frame)',
          tsize=14, bsize=11, tcol=TCOL["if"], title_top=True)

    # VLA 子模块
    rrect(slide, 72.5, 70, 23, 11, "FFFFFF", COL["vla"][1], "PaliGemma\nVision-Language Trunk", tsize=12)
    rrect(slide, 72.5, 56, 23, 11, "FFFFFF", COL["vla"][1], "Flow-Matching\nAction Expert", tsize=12)
    rrect(slide, 72.5, 43, 23, 10, "EAF6EE", COL["vla"][1],
          "Action sequence (next K steps)\n6-DoF Δpose / joint targets", tsize=11, bold=False)
    arrow(slide, 84, 70, 84, 67, COL["vla"][1]); arrow(slide, 84, 56, 84, 53, COL["vla"][1])

    # 音频闭环
    rrect(slide, 70, 10, 28, 19, *COL["aud"], "③  Audio Loop   (post-contact)",
          body="audio quality → search (I,U,V)\nneighborhood → refine target",
          tsize=14, bsize=11, tcol=TCOL["aud"], title_top=True)

    # 机器人
    rrect(slide, 30, 11, 26, 13, *COL["rob"], "Robot arm\n+ stethoscope", tsize=13, tcol=TCOL["rob"])

    # 连线
    arrow(slide, 18, 79, 23, 74, COL["per"][1]); textlabel(slide, 21, 81, "RGB", col=COL["per"][1], w=8)
    arrow(slide, 44, 58.6, 44, 54.5, COL["per"][1])
    arrow(slide, 65, 49, 70, 58, COL["if"][1], width=2.4)
    textlabel(slide, 68.5, 54, "heatmap + 3D point", col="A07410", size=9, w=18)
    arrow(slide, 18, 42, 70, 50, COL["vla"][1])
    textlabel(slide, 44, 37, "wrist RGB · proprioception · force/torque · audio · language",
              col=COL["vla"][1], size=9.5, w=58)
    arrow(slide, 72.5, 44, 56, 19, COL["vla"][1]); textlabel(slide, 66, 31, "actions", col=COL["vla"][1], w=10)
    arrow(slide, 56, 15, 70, 16, COL["aud"][1]); textlabel(slide, 63, 12.5, "contact", col=COL["aud"][1], size=9, w=10)
    arrow(slide, 80, 29, 60, 45, COL["aud"][1], dashed=True)
    textlabel(slide, 72, 38, "refine (I,U,V)", col=COL["aud"][1], size=9, w=14)

    # 脚注
    textlabel(slide, 50, 4,
              "Modules trained separately, coupled only through the target hand-off interface.   "
              "①  Perception: dense-IUV supervision on body-image data (HuMMan / Scoliosis / Synthetic).   "
              "②  VLA: behavior cloning on teleoperation demos.",
              size=9, col="444444", w=95, h=5)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    prs.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
